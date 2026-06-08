"""Excel-Generator (ALE-33) — ROI-Rechner + Tool-Vergleichsmatrix aus FullReport.

Erzeugt eine .xlsx-Arbeitsmappe mit zwei Arbeitsblättern:
  1. "ROI-Rechner"      — pro Use-Case Investment + Savings Y1-3 + Payback + 3J-ROI,
                          abschließende Summenzeile, bedingte Hervorhebung der
                          besten (kürzesten) Payback-Zeile.
  2. "Tool-Vergleich"   — Vergleichsmatrix aller Tool-Empfehlungen
                          (Vendor, Kosten, Setup, Time-to-Value, Compliance).

Alle Zahlen stammen ausschließlich aus den Pydantic-Objekten des FullReport —
keine erfundenen Werte. Brand-Farben gemäß Design-System (Navy/Magenta/Teal/Amber).

Öffentliche API:
    build_excel(report: FullReport) -> bytes
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from schemas.outputs import FullReport, ROIOutput, ToolRecommendationOutput

# — Brand-Farben (ohne führendes #, wie openpyxl sie erwartet) —
NAVY = "0E2244"
MAGENTA = "E94B5D"
TEAL = "1AA591"
AMBER = "F59E0B"
CREAM = "FDF8F0"
WHITE = "FFFFFF"
SLATE = "475569"
LIGHT_ROW = "F4F6F9"

EUR_FMT = '#,##0 "€"'
PCT_FMT = '0 "%"'
MONTHS_FMT = '0.0 "Mo"'
HOURS_FMT = '0.0 "h"'

_THIN = Side(style="thin", color="D5DAE2")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_fill(color: str = NAVY) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _white_bold() -> Font:
    return Font(name="Lato", bold=True, color=WHITE, size=11)


def _body_font(bold: bool = False, color: str = NAVY) -> Font:
    return Font(name="Lato", bold=bold, color=color, size=10)


def _style_header_row(ws: Worksheet, row: int, n_cols: int, color: str = NAVY) -> None:
    fill = _header_fill(color)
    font = _white_bold()
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 34


def _title_block(ws: Worksheet, title: str, subtitle: str, n_cols: int) -> int:
    """Schreibt Titel (Zeile 1) + Untertitel (Zeile 2), gibt nächste freie Zeile zurück."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    tcell = ws.cell(row=1, column=1, value=title)
    tcell.font = Font(name="Playfair Display", bold=True, color=NAVY, size=16)
    tcell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    scell = ws.cell(row=2, column=2 if False else 1, value=subtitle)
    scell.font = Font(name="Lato", italic=True, color=SLATE, size=10)
    scell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20
    return 4  # Zeile 3 bleibt Luft, Header beginnt in Zeile 4


def _set_widths(ws: Worksheet, widths: list[int]) -> None:
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 1 — ROI-Rechner
# ─────────────────────────────────────────────────────────────────────────────
def _build_roi_sheet(ws: Worksheet, roi: ROIOutput, company_name: str) -> None:
    headers = [
        "Use-Case", "Prozess", "Investment Y1 (€)", "Setup (€)",
        "Savings Y1 (€)", "Savings Y2 (€)", "Savings Y3 (€)",
        "Zeitersparnis (h/Wo)", "Payback (Mo)", "3J-ROI (%)", "Konfidenz",
    ]
    n_cols = len(headers)
    header_row = _title_block(
        ws,
        f"ROI-Rechner · {company_name}",
        "Investition vs. Einsparung pro Use-Case — Zahlen aus dem ROI-Calculator-Agent.",
        n_cols,
    )

    for col, label in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=label)
    _style_header_row(ws, header_row, n_cols)

    # — beste (kürzeste) Payback-Zeile ermitteln (nur unter Items mit Savings > 0) —
    best_idx = None
    best_payback = None
    for i, li in enumerate(roi.line_items):
        if li.payback_months and li.payback_months > 0:
            if best_payback is None or li.payback_months < best_payback:
                best_payback = li.payback_months
                best_idx = i

    first_data_row = header_row + 1
    highlight_fill = _header_fill(TEAL)
    alt_fill = _header_fill(LIGHT_ROW)

    for i, li in enumerate(roi.line_items):
        row = first_data_row + i
        is_best = i == best_idx
        values = [
            li.use_case_name,
            li.target_process,
            li.investment_eur_year_1,
            li.setup_cost_eur,
            li.savings_eur_year_1,
            li.savings_eur_year_2,
            li.savings_eur_year_3,
            li.time_saved_h_week,
            li.payback_months,
            li.three_year_roi_percent,
            li.confidence,
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="left" if col <= 2 else "right",
                vertical="center",
            )
            if is_best:
                cell.font = _body_font(bold=True, color=WHITE)
                cell.fill = highlight_fill
            else:
                cell.font = _body_font()
                if i % 2 == 1:
                    cell.fill = alt_fill
        # Zahlenformate
        for col in (3, 4, 5, 6, 7):
            ws.cell(row=row, column=col).number_format = EUR_FMT
        ws.cell(row=row, column=8).number_format = HOURS_FMT
        ws.cell(row=row, column=9).number_format = MONTHS_FMT
        ws.cell(row=row, column=10).number_format = PCT_FMT
        ws.cell(row=row, column=11).alignment = Alignment(horizontal="center", vertical="center")

    # — Summenzeile (Werte aus ROIOutput-Aggregaten, nicht neu berechnet) —
    sum_row = first_data_row + len(roi.line_items)
    sum_fill = _header_fill(NAVY)
    sum_values = {
        1: "SUMME",
        3: roi.total_investment_eur,
        4: roi.total_setup_cost_eur,
        5: roi.total_savings_eur_year_1,
        6: roi.total_savings_eur_year_2,
        7: roi.total_savings_eur_year_3,
        9: roi.total_payback_months,
    }
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=sum_row, column=col, value=sum_values.get(col))
        cell.fill = sum_fill
        cell.font = _white_bold()
        cell.border = BORDER
        cell.alignment = Alignment(
            horizontal="left" if col <= 2 else "right", vertical="center"
        )
    for col in (3, 4, 5, 6, 7):
        ws.cell(row=sum_row, column=col).number_format = EUR_FMT
    ws.cell(row=sum_row, column=9).number_format = MONTHS_FMT
    ws.row_dimensions[sum_row].height = 22

    # — 3-Jahres-Savings als eigene Amber-Hervorhebung (max. 1× Amber pro Sheet) —
    kpi_row = sum_row + 2
    ws.cell(row=kpi_row, column=1, value="3-Jahres-Einsparung gesamt").font = _body_font(bold=True)
    kpi_cell = ws.cell(row=kpi_row, column=3, value=roi.total_savings_eur_3_years)
    kpi_cell.number_format = EUR_FMT
    kpi_cell.fill = _header_fill(AMBER)
    kpi_cell.font = Font(name="Lato", bold=True, color=NAVY, size=11)
    kpi_cell.alignment = Alignment(horizontal="right", vertical="center")
    kpi_cell.border = BORDER

    # — Sensitivitäts-Hinweis + Summary —
    note_row = kpi_row + 2
    if roi.sensitivity_notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=n_cols)
        c = ws.cell(row=note_row, column=1, value=f"Sensitivität: {roi.sensitivity_notes}")
        c.font = Font(name="Lato", italic=True, color=SLATE, size=9)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[note_row].height = 40
        note_row += 1
    if roi.summary:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=n_cols)
        c = ws.cell(row=note_row, column=1, value=f"Fazit: {roi.summary}")
        c.font = Font(name="Lato", color=SLATE, size=9)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[note_row].height = 48

    _set_widths(ws, [26, 22, 16, 12, 14, 14, 14, 16, 12, 12, 12])
    ws.freeze_panes = ws.cell(row=first_data_row, column=1)


# ─────────────────────────────────────────────────────────────────────────────
# Sheet 2 — Tool-Vergleichsmatrix
# ─────────────────────────────────────────────────────────────────────────────
def _yesno(value: bool | None) -> str:
    if value is None:
        return "?"
    return "ja" if value else "nein"


def _build_tools_sheet(ws: Worksheet, tools: ToolRecommendationOutput, company_name: str) -> None:
    headers = [
        "Use-Case", "Empfehlung (Tool)", "Land", "Preismodell",
        "€/Monat", "Setup (€)", "Setup-Komplexität", "Time-to-Value (Wo)",
        "EU-Hosting", "AVV", "AI-Act-Klasse", "Alternativen",
    ]
    n_cols = len(headers)
    header_row = _title_block(
        ws,
        f"Tool-Vergleichsmatrix · {company_name}",
        "Empfohlener Stack aus dem Tool-Recommender-Agent — Kosten, Setup, Compliance.",
        n_cols,
    )

    for col, label in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=label)
    _style_header_row(ws, header_row, n_cols, color=NAVY)

    # — günstigste monatliche Empfehlung hervorheben (Teal) —
    best_idx = None
    best_cost = None
    for i, rec in enumerate(tools.recommendations):
        if best_cost is None or rec.monthly_cost_eur < best_cost:
            best_cost = rec.monthly_cost_eur
            best_idx = i

    first_data_row = header_row + 1
    highlight_fill = _header_fill(TEAL)
    alt_fill = _header_fill(LIGHT_ROW)

    for i, rec in enumerate(tools.recommendations):
        row = first_data_row + i
        is_best = i == best_idx
        values = [
            rec.use_case_name,
            rec.primary_tool,
            rec.primary_tool_vendor_country,
            rec.primary_tool_pricing_model,
            rec.monthly_cost_eur,
            rec.setup_cost_eur,
            rec.setup_complexity,
            rec.time_to_value_weeks,
            _yesno(rec.compliance.eu_hosting),
            _yesno(rec.compliance.avv_available),
            rec.compliance.ai_act_class,
            ", ".join(rec.alternative_tools) if rec.alternative_tools else "—",
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="right" if col in (5, 6, 8) else "left",
                vertical="center",
                wrap_text=col in (4, 12),
            )
            if is_best:
                cell.font = _body_font(bold=True, color=WHITE)
                cell.fill = highlight_fill
            else:
                cell.font = _body_font()
                if i % 2 == 1:
                    cell.fill = alt_fill
        ws.cell(row=row, column=5).number_format = EUR_FMT
        ws.cell(row=row, column=6).number_format = EUR_FMT
        ws.cell(row=row, column=8).number_format = '0 "Wo"'

    # — Summenzeile (aus ToolRecommendationOutput-Aggregaten) —
    sum_row = first_data_row + len(tools.recommendations)
    sum_fill = _header_fill(NAVY)
    sum_values = {
        1: "SUMME STACK",
        5: tools.total_monthly_cost_eur,
        6: tools.total_setup_cost_eur,
    }
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=sum_row, column=col, value=sum_values.get(col))
        cell.fill = sum_fill
        cell.font = _white_bold()
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="right" if col in (5, 6) else "left", vertical="center")
    ws.cell(row=sum_row, column=5).number_format = EUR_FMT
    ws.cell(row=sum_row, column=6).number_format = EUR_FMT
    ws.row_dimensions[sum_row].height = 22

    # — Stack-Summary —
    if tools.stack_summary:
        note_row = sum_row + 2
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=n_cols)
        c = ws.cell(row=note_row, column=1, value=f"Stack-Logik: {tools.stack_summary}")
        c.font = Font(name="Lato", color=SLATE, size=9)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[note_row].height = 48

    _set_widths(ws, [22, 24, 10, 18, 12, 12, 16, 16, 12, 8, 14, 28])
    ws.freeze_panes = ws.cell(row=first_data_row, column=1)


# ─────────────────────────────────────────────────────────────────────────────
# Öffentliche API
# ─────────────────────────────────────────────────────────────────────────────
def build_excel(report: FullReport) -> bytes:
    """Baut die ROI-/Tool-Arbeitsmappe und gibt die .xlsx als Bytes zurück."""
    wb = Workbook()

    roi_ws = wb.active
    roi_ws.title = "ROI-Rechner"
    _build_roi_sheet(roi_ws, report.roi, report.company_name)

    tools_ws = wb.create_sheet("Tool-Vergleich")
    _build_tools_sheet(tools_ws, report.tools, report.company_name)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

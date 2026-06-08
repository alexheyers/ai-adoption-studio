"""PPTX-Generator (ALE-32) — 10-Slide-Board-Deck aus FullReport via python-pptx.

Slide-Architektur (10 Slides, 16:9):
   1  Titel                — Firma + Executive-Summary-Teaser
   2  Executive Summary    — die 3-5-Sätze-Empfehlung des Reporters
   3  Ausgangslage         — Audit-Summary + Domain-Abdeckung
   4  Prozess-Hotspots     — Top-Prozesse mit Automatisierungs-Potenzial
   5  Use-Cases            — Use-Case-Liste + Quick-Wins
   6  Empfohlener Stack    — Tool-Empfehlungen (Tabelle)
   7  ROI                  — Kern-Zahlen Investment/Savings/Payback (Tabelle)
   8  Compliance           — DSGVO/AI-Act-Flags (optional, sonst Hinweis)
   9  Roadmap              — 3 Phasen
  10  Next Steps           — Handlungsempfehlung + Critical-Path

Alle Inhalte werden aus dem FullReport gebunden — keine erfundenen Zahlen.
Brand-Farben: Navy #0E2244, Magenta #E94B5D, Teal #1AA591, Amber #F59E0B
(Amber max. 1× pro Slide). Fonts Lato/Playfair, Fallback Calibri.

Öffentliche API:
    build_pptx(report: FullReport) -> bytes
"""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from schemas.outputs import FullReport

# — Brand-Farben —
NAVY = RGBColor(0x0E, 0x22, 0x44)
MAGENTA = RGBColor(0xE9, 0x4B, 0x5D)
TEAL = RGBColor(0x1A, 0xA5, 0x91)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
CREAM = RGBColor(0xFD, 0xF8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0xF4, 0xF6, 0xF9)

HEAD_FONT = "Playfair Display"
BODY_FONT = "Lato"

# 16:9-Standardmaße
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _blank(prs: Presentation):
    """Leeres Layout (i.d.R. Index 6 im Default-Template)."""
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    return prs.slides.add_slide(layout)


def _bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, left, top, width, height, color: RGBColor):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _text(
    slide, left, top, width, height, text,
    *, size=18, color=NAVY, bold=False, italic=False,
    font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text or ""
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    return box


def _bullets(
    slide, left, top, width, height, items,
    *, size=14, color=NAVY, font=BODY_FONT, bullet_color=TEAL, gap=4,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        marker = p.add_run()
        marker.text = "▸  "
        marker.font.size = Pt(size)
        marker.font.name = font
        marker.font.bold = True
        marker.font.color.rgb = bullet_color
        run = p.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = color
    return box


def _accent_bar(slide, color: RGBColor = TEAL) -> None:
    _rect(slide, Inches(0), Inches(0), Inches(0.12), SLIDE_H, color)


def _section_header(slide, kicker: str, title: str, *, accent: RGBColor = TEAL):
    """Standard-Kopf für Content-Slides auf hellem Grund."""
    _bg(slide, CREAM)
    _accent_bar(slide, accent)
    _text(slide, Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.4), kicker.upper(),
          size=12, color=accent, bold=True, font=BODY_FONT)
    _text(slide, Inches(0.7), Inches(0.85), Inches(11.9), Inches(0.9), title,
          size=30, color=NAVY, bold=True, font=HEAD_FONT)
    _rect(slide, Inches(0.72), Inches(1.72), Inches(1.3), Inches(0.05), accent)


def _table(slide, left, top, width, rows, cols, col_widths=None):
    gfx = slide.shapes.add_table(rows, cols, left, top, width, Inches(0.4 * rows))
    table = gfx.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    return table


def _style_cell(cell, text, *, size=11, color=NAVY, bold=False,
                fill: RGBColor | None = None, align=PP_ALIGN.LEFT, font=BODY_FONT):
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color


def _fmt_eur(value: int | float | None) -> str:
    if value is None:
        return "—"
    return f"{int(round(value)):,} €".replace(",", ".")


# ─────────────────────────────────────────────────────────────────────────────
# Einzelne Slides
# ─────────────────────────────────────────────────────────────────────────────
def _slide_title(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _bg(slide, NAVY)
    _rect(slide, Inches(0), Inches(0), Inches(0.18), SLIDE_H, TEAL)
    _text(slide, Inches(0.9), Inches(1.6), Inches(11.5), Inches(0.5),
          "AI-ADOPTION-STUDIO · BOARD-REPORT",
          size=14, color=TEAL, bold=True, font=BODY_FONT)
    _text(slide, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.0),
          report.company_name, size=48, color=WHITE, bold=True, font=HEAD_FONT)
    _rect(slide, Inches(0.95), Inches(3.7), Inches(2.2), Inches(0.08), AMBER)
    teaser = (report.executive_summary or "").strip()
    if len(teaser) > 220:
        teaser = teaser[:217].rstrip() + "…"
    _text(slide, Inches(0.9), Inches(4.0), Inches(11.0), Inches(2.5),
          teaser, size=18, color=CREAM, italic=True, font=BODY_FONT)


def _slide_exec_summary(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _bg(slide, NAVY)
    _accent_bar(slide, TEAL)
    _text(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.4),
          "EXECUTIVE SUMMARY", size=13, color=TEAL, bold=True)
    _text(slide, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.9),
          "Die Empfehlung in 30 Sekunden", size=30, color=WHITE, bold=True, font=HEAD_FONT)
    _text(slide, Inches(0.7), Inches(2.0), Inches(11.9), Inches(4.5),
          report.executive_summary or "—", size=20, color=CREAM, font=BODY_FONT)


def _slide_situation(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _section_header(slide, "Ausgangslage", "Wo das Haus heute steht")
    _text(slide, Inches(0.7), Inches(2.0), Inches(11.9), Inches(2.8),
          report.audit.summary or "—", size=16, color=SLATE, font=BODY_FONT)
    if report.audit.domains_unchecked:
        _text(slide, Inches(0.7), Inches(4.9), Inches(11.9), Inches(0.4),
              "Nicht beurteilbar (Datenlücken):", size=13, color=NAVY, bold=True)
        _bullets(slide, Inches(0.7), Inches(5.3), Inches(11.9), Inches(1.6),
                 report.audit.domains_unchecked[:6], size=13, color=SLATE, bullet_color=MAGENTA)


def _slide_hotspots(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _section_header(slide, "Prozess-Hotspots", "Die größten Hebel")
    rank = {"high": 0, "medium": 1, "low": 2}
    procs = sorted(
        report.audit.processes,
        key=lambda p: (rank.get(p.automation_potential, 3), -(p.estimated_savings_eur_year or 0)),
    )[:5]
    if not procs:
        _text(slide, Inches(0.7), Inches(2.2), Inches(11), Inches(1), "Keine Prozesse erfasst.", size=16, color=SLATE)
        return
    rows = len(procs) + 1
    table = _table(slide, Inches(0.7), Inches(2.0), Inches(11.9), rows, 4,
                   col_widths=[Inches(4.1), Inches(2.0), Inches(2.6), Inches(3.2)])
    for c, h in enumerate(["Prozess", "Potenzial", "Einsparung/Jahr", "Hauptschmerz"]):
        _style_cell(table.cell(0, c), h, size=12, color=WHITE, bold=True, fill=NAVY)
    for i, p in enumerate(procs, start=1):
        fill = WHITE if i % 2 else LIGHT
        _style_cell(table.cell(i, 0), p.name, size=11, bold=True, fill=fill)
        pot = p.automation_potential
        pot_color = TEAL if pot == "high" else (NAVY if pot == "medium" else SLATE)
        _style_cell(table.cell(i, 1), pot, size=11, color=pot_color, bold=(pot == "high"), fill=fill, align=PP_ALIGN.CENTER)
        _style_cell(table.cell(i, 2), _fmt_eur(p.estimated_savings_eur_year), size=11, fill=fill, align=PP_ALIGN.RIGHT)
        pain = (p.primary_pain or "")[:90]
        _style_cell(table.cell(i, 3), pain, size=10, color=SLATE, fill=fill)


def _slide_use_cases(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _section_header(slide, "Use-Cases", "Was wir konkret bauen würden")
    ucs = report.use_cases.use_cases
    quick = [u for u in ucs if u.quick_win]
    rest = [u for u in ucs if not u.quick_win]
    ordered = (quick + rest)[:6]
    items = []
    for u in ordered:
        tag = "Quick-Win · " if u.quick_win else ""
        items.append(f"{tag}{u.name} — {u.expected_impact[:70]}")
    if items:
        _bullets(slide, Inches(0.7), Inches(2.0), Inches(11.9), Inches(3.6),
                 items, size=15, color=NAVY, gap=8)
    # Quick-Wins-Summary als Amber-Kasten (einzige Amber-Fläche auf der Slide)
    if report.use_cases.quick_wins_summary:
        _rect(slide, Inches(0.7), Inches(5.7), Inches(11.9), Inches(1.3), AMBER)
        _text(slide, Inches(0.95), Inches(5.85), Inches(11.4), Inches(1.0),
              f"Quick-Wins: {report.use_cases.quick_wins_summary[:200]}",
              size=13, color=NAVY, bold=True, font=BODY_FONT)


def _slide_stack(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _section_header(slide, "Empfohlener Stack", "Tools, die wir empfehlen")
    recs = report.tools.recommendations[:5]
    if not recs:
        _text(slide, Inches(0.7), Inches(2.2), Inches(11), Inches(1), "Keine Tools empfohlen.", size=16, color=SLATE)
        return
    rows = len(recs) + 2  # + Header + Summe
    table = _table(slide, Inches(0.7), Inches(2.0), Inches(11.9), rows, 4,
                   col_widths=[Inches(3.6), Inches(3.9), Inches(2.2), Inches(2.2)])
    for c, h in enumerate(["Use-Case", "Tool", "€/Monat", "Setup"]):
        _style_cell(table.cell(0, c), h, size=12, color=WHITE, bold=True, fill=NAVY)
    for i, r in enumerate(recs, start=1):
        fill = WHITE if i % 2 else LIGHT
        _style_cell(table.cell(i, 0), (r.use_case_name or "")[:40], size=11, fill=fill)
        _style_cell(table.cell(i, 1), (r.primary_tool or "")[:48], size=11, bold=True, fill=fill)
        _style_cell(table.cell(i, 2), _fmt_eur(r.monthly_cost_eur), size=11, fill=fill, align=PP_ALIGN.RIGHT)
        _style_cell(table.cell(i, 3), _fmt_eur(r.setup_cost_eur), size=11, fill=fill, align=PP_ALIGN.RIGHT)
    last = rows - 1
    _style_cell(table.cell(last, 0), "Stack gesamt", size=12, color=WHITE, bold=True, fill=TEAL)
    _style_cell(table.cell(last, 1), "", fill=TEAL)
    _style_cell(table.cell(last, 2), _fmt_eur(report.tools.total_monthly_cost_eur), size=12, color=WHITE, bold=True, fill=TEAL, align=PP_ALIGN.RIGHT)
    _style_cell(table.cell(last, 3), _fmt_eur(report.tools.total_setup_cost_eur), size=12, color=WHITE, bold=True, fill=TEAL, align=PP_ALIGN.RIGHT)


def _slide_roi(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _section_header(slide, "ROI", "Was es kostet, was es bringt")
    roi = report.roi
    # KPI-Kacheln
    kpis = [
        ("Investment Jahr 1", _fmt_eur(roi.total_investment_eur), NAVY),
        ("Einsparung Jahr 1", _fmt_eur(roi.total_savings_eur_year_1), TEAL),
        ("Einsparung 3 Jahre", _fmt_eur(roi.total_savings_eur_3_years), TEAL),
        ("Payback", f"{roi.total_payback_months:.1f} Mo", AMBER),
    ]
    x = Inches(0.7)
    w = Inches(2.85)
    gap = Inches(0.13)
    for label, val, col in kpis:
        _rect(slide, x, Inches(2.0), w, Inches(1.5), col)
        text_col = NAVY if col == AMBER else WHITE
        _text(slide, x, Inches(2.15), w, Inches(0.5), label, size=12, color=text_col, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, x, Inches(2.65), w, Inches(0.7), val, size=22, color=text_col, bold=True, font=HEAD_FONT, align=PP_ALIGN.CENTER)
        x = Emu(x + w + gap)
    # Line-Items-Tabelle (Top 4)
    items = report.roi.line_items[:4]
    rows = len(items) + 1
    if items:
        table = _table(slide, Inches(0.7), Inches(3.9), Inches(11.9), rows, 4,
                       col_widths=[Inches(4.7), Inches(2.4), Inches(2.4), Inches(2.4)])
        for c, h in enumerate(["Use-Case", "Investment Y1", "Savings Y1", "Payback"]):
            _style_cell(table.cell(0, c), h, size=11, color=WHITE, bold=True, fill=NAVY)
        for i, li in enumerate(items, start=1):
            fill = WHITE if i % 2 else LIGHT
            _style_cell(table.cell(i, 0), (li.use_case_name or "")[:48], size=10, fill=fill)
            _style_cell(table.cell(i, 1), _fmt_eur(li.investment_eur_year_1), size=10, fill=fill, align=PP_ALIGN.RIGHT)
            _style_cell(table.cell(i, 2), _fmt_eur(li.savings_eur_year_1), size=10, fill=fill, align=PP_ALIGN.RIGHT)
            _style_cell(table.cell(i, 3), f"{li.payback_months} Mo", size=10, fill=fill, align=PP_ALIGN.RIGHT)


def _slide_compliance(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _section_header(slide, "Compliance", "DSGVO & AI-Act im Blick", accent=MAGENTA)
    comp = report.compliance
    if not comp or not comp.flags:
        _text(slide, Inches(0.7), Inches(2.2), Inches(11.5), Inches(1.5),
              "Keine kritischen Compliance-Flags im aktuellen Scope erfasst. "
              "Vor Umsetzung dennoch mit DSGVO-Beauftragten prüfen.",
              size=16, color=SLATE, font=BODY_FONT)
        return
    flagged = [f for f in comp.flags if f.dsgvo_relevant or f.ai_act_risk_class != "minimal"]
    show = (flagged or comp.flags)[:5]
    items = []
    for f in show:
        cls = f.ai_act_risk_class
        dsgvo = "DSGVO" if f.dsgvo_relevant else "—"
        items.append(f"{f.use_case_name} · AI-Act: {cls} · {dsgvo}")
    _bullets(slide, Inches(0.7), Inches(2.0), Inches(11.9), Inches(3.0),
             items, size=15, color=NAVY, bullet_color=MAGENTA, gap=8)
    if comp.general_advice:
        _text(slide, Inches(0.7), Inches(5.2), Inches(11.9), Inches(1.6),
              comp.general_advice[:280], size=12, color=SLATE, italic=True)


def _slide_roadmap(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _section_header(slide, "Roadmap", "Drei Phasen bis zur Adoption")
    rm = report.roadmap
    if not rm or not rm.phases:
        _text(slide, Inches(0.7), Inches(2.2), Inches(11), Inches(1), "Keine Roadmap erstellt.", size=16, color=SLATE)
        return
    phases = sorted(rm.phases, key=lambda p: p.phase_number)[:3]
    colors = [TEAL, NAVY, MAGENTA]
    x = Inches(0.7)
    w = Inches(3.85)
    gap = Inches(0.2)
    for idx, ph in enumerate(phases):
        col = colors[idx % len(colors)]
        _rect(slide, x, Inches(2.0), w, Inches(0.7), col)
        _text(slide, x, Inches(2.08), w, Inches(0.55),
              f"Phase {ph.phase_number} · {ph.duration_months}",
              size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, Emu(x + Inches(0.1)), Inches(2.85), Emu(w - Inches(0.2)), Inches(0.6),
              ph.name, size=15, color=NAVY, bold=True, font=HEAD_FONT)
        body = list(ph.use_cases[:3]) + [f"Aufwand: {ph.estimated_effort_pt} PT"]
        _bullets(slide, Emu(x + Inches(0.1)), Inches(3.5), Emu(w - Inches(0.2)), Inches(2.8),
                 body, size=11, color=SLATE, bullet_color=col, gap=5)
        x = Emu(x + w + gap)
    if rm.critical_path:
        _text(slide, Inches(0.7), Inches(6.4), Inches(11.9), Inches(0.8),
              f"Critical Path: {rm.critical_path[:160]}", size=12, color=NAVY, italic=True, bold=True)


def _slide_next_steps(prs, report: FullReport) -> None:
    slide = _blank(prs)
    _bg(slide, NAVY)
    _accent_bar(slide, TEAL)
    _text(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.4),
          "NÄCHSTE SCHRITTE", size=13, color=TEAL, bold=True)
    _text(slide, Inches(0.7), Inches(0.95), Inches(11.9), Inches(0.9),
          "Womit wir morgen starten", size=32, color=WHITE, bold=True, font=HEAD_FONT)

    steps: list[str] = []
    quick = [u.name for u in report.use_cases.use_cases if u.quick_win][:2]
    for q in quick:
        steps.append(f"Quick-Win starten: {q}")
    if report.roadmap and report.roadmap.phases:
        p1 = sorted(report.roadmap.phases, key=lambda p: p.phase_number)[0]
        steps.append(f"Phase 1 ({p1.duration_months}): {p1.name}")
    steps.append(
        f"Business-Case bestätigen: Payback {report.roi.total_payback_months:.1f} Monate, "
        f"3J-Einsparung {_fmt_eur(report.roi.total_savings_eur_3_years)}"
    )
    if report.roadmap and report.roadmap.critical_path:
        steps.append(f"Engpass früh adressieren: {report.roadmap.critical_path[:90]}")

    _bullets(slide, Inches(0.7), Inches(2.2), Inches(11.9), Inches(3.6),
             steps, size=18, color=CREAM, bullet_color=TEAL, gap=12)

    # Amber-Schluss-Akzent (einzige Amber-Fläche)
    _rect(slide, Inches(0.7), Inches(6.4), Inches(4.0), Inches(0.08), AMBER)
    _text(slide, Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.6),
          "AI-Adoption-Studio · datengestützte Empfehlung, kein Marketing.",
          size=12, color=SLATE, italic=True)


# ─────────────────────────────────────────────────────────────────────────────
# Öffentliche API
# ─────────────────────────────────────────────────────────────────────────────
def build_pptx(report: FullReport) -> bytes:
    """Baut das 10-Slide-Deck und gibt die .pptx als Bytes zurück."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _slide_title(prs, report)
    _slide_exec_summary(prs, report)
    _slide_situation(prs, report)
    _slide_hotspots(prs, report)
    _slide_use_cases(prs, report)
    _slide_stack(prs, report)
    _slide_roi(prs, report)
    _slide_compliance(prs, report)
    _slide_roadmap(prs, report)
    _slide_next_steps(prs, report)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
